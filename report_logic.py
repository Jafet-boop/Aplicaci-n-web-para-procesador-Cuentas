import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import os
from dbfread import DBF

# -----------------------------
# 1. Carga y limpieza
# -----------------------------
def load_and_clean(path: str) -> pd.DataFrame:
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        df = pd.read_csv(path)
    elif ext in [".xls", ".xlsx"]:
        df = pd.read_excel(path)
    elif ext == ".dbf":
        table = DBF(path, encoding="latin1", ignore_missing_memofile=True)
        df = pd.DataFrame(iter(table))
    else:
        raise ValueError(f"Formato no soportado: {ext}")

    df = df.rename(columns={
        "FECHAVTO": "FECHAVTO",
        "NOFAC": "NOFAC",
        "PROVEEDOR": "PROVEEDOR",
        " SALDO ": "SALDO",
        "SALDO": "SALDO"
    })

    def clean_saldo(x):
        if pd.isna(x):
            return None
        x = str(x).replace("$", "").replace(",", "").strip()
        if x in ["-", "", "$-", " -"]:
            return None
        try:
            return float(x)
        except:
            return None

    df["SALDO"] = df["SALDO"].apply(clean_saldo)
    df = df[df["SALDO"].notna()]

    df["FECHAVTO"] = pd.to_datetime(df["FECHAVTO"], format="%d/%m/%Y")

    meses = {
        1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
        5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
        9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"
    }

    df["MES"] = df["FECHAVTO"].dt.month.map(meses)
    df["AÑO"] = df["FECHAVTO"].dt.year

    hoy = pd.Timestamp.today().normalize()
    df["ESTATUS"] = df["FECHAVTO"].apply(
        lambda x: "VENCIDO" if x < hoy else "POR VENCER"
    )

    df = df.sort_values(by=["PROVEEDOR", "FECHAVTO"])

    return df

# -----------------------------
# 2. Generar Excel analítico
# -----------------------------
def generate_xlsx(df: pd.DataFrame, output_path: str):
    wb = openpyxl.Workbook()

    bold = Font(bold=True)
    bold_big = Font(bold=True, size=14)
    right = Alignment(horizontal="right")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

# ======================================================
# HOJA 1 – KPIs
# ======================================================
    ws1 = wb.active
    ws1.title = "KPIs"

    header_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    row_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    ws1.cell(row=1, column=1, value="KPIs – Cuentas por Pagar").font = bold_big

    # Encabezados para la tabla
    ws1.cell(row=2, column=1, value="Métrica").font = bold
    ws1.cell(row=2, column=2, value="Valor").font = bold

    ws1.cell(row=2, column=1).fill = header_fill
    ws1.cell(row=2, column=2).fill = header_fill

    total_general = df["SALDO"].sum()
    total_vencido = df[df["ESTATUS"] == "VENCIDO"]["SALDO"].sum()
    total_por_vencer = df[df["ESTATUS"] == "POR VENCER"]["SALDO"].sum()
    num_proveedores = df["PROVEEDOR"].nunique()

    porcentaje_vencido = total_vencido / total_general if total_general > 0 else 0


    kpis = [
        ("Total general por pagar", total_general),
        ("Total vencido", total_vencido),
        ("Total por vencer", total_por_vencer),
        ("% vencido", porcentaje_vencido),
        ("Número de proveedores", num_proveedores),
    ]

    row = 3
    for label, value in kpis:
        ws1.cell(row=row, column=1, value=label)
        c = ws1.cell(row=row, column=2, value=value)
        c.alignment = right

        # Formato
        ws1.cell(row=row, column=1).fill = row_fill
        ws1.cell(row=row, column=2).fill = row_fill

        ws1.cell(row=row, column=1).border = border
        ws1.cell(row=row, column=2).border = border

        row += 1

    # Bordes encabezados
        ws1.cell(row=2, column=1).border = border
        ws1.cell(row=2, column=2).border = border

        # Ancho de columnas
        ws1.column_dimensions["A"].width = 40
        ws1.column_dimensions["B"].width = 20
# ======================================================
# HOJA 2 – CALENDARIO DE PAGOS
# ======================================================
    ws2 = wb.create_sheet(title="Calendario de Pagos")
    ws2.cell(row=1, column=1, value="Calendario de Pagos").font = bold_big
    headers = ["FECHA VENCIMIENTO", "PROVEEDOR", "FACTURA", "SALDO", "ESTATUS"]
    row = 3

    # Colores
    fill_header = PatternFill("solid", fgColor="CFE2F3")   # azul claro
    fill_body = PatternFill("solid", fgColor="F9F9F9")     # gris claro
    fill_vencido = PatternFill("solid", fgColor="F4CCCC") # rojo suave
    fill_por_vencer = PatternFill("solid", fgColor="D9EAD3") # verde suave

    # Encabezados
    for col, h in enumerate(headers, start=1):
        cell = ws2.cell(row=row, column=col, value=h)
        cell.font = bold
        cell.fill = fill_header
        cell.border = border
    row += 1

    # Datos
    calendario = df.sort_values(by="FECHAVTO")
    data_start = row

    for _, r in calendario.iterrows():
        ws2.cell(row=row, column=1, value=r["FECHAVTO"].strftime("%d/%m/%Y"))
        ws2.cell(row=row, column=2, value=r["PROVEEDOR"])
        ws2.cell(row=row, column=3, value=r["NOFAC"])

        c_saldo = ws2.cell(row=row, column=4, value=r["SALDO"])
        c_saldo.alignment = right

        c_status = ws2.cell(row=row, column=5, value=r["ESTATUS"])
        
        for c in range(1, 5):  # ← Cambié de 6 a 5 para excluir la columna de estatus
           cell = ws2.cell(row=row, column=c)
           cell.fill = fill_body
           cell.border = border
    
    # Aplicar borde y color solo a la columna de estatus
        c_status.border = border
        if r["ESTATUS"] == "VENCIDO":
            c_status.fill = fill_vencido
        else:
            c_status.fill = fill_por_vencer
        
        row += 1

        # Ancho de columnas
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 35
    ws2.column_dimensions["C"].width = 20
    ws2.column_dimensions["D"].width = 18
    ws2.column_dimensions["E"].width = 18

# ======================================================
# HOJA 3 – MATRIZ GENERAL POR PROVEEDOR Y MES
# ======================================================
    ws3 = wb.create_sheet(title="Matriz")
    ws3.cell(row=1, column=1, value="Matriz de Saldos por Proveedor y Mes").font = bold_big

    # Colores
    fill_header = PatternFill("solid", fgColor="54F23F")   # verde claro
    fill_body = PatternFill("solid", fgColor="F6FBF4")     # verde muy suave
    fill_total = PatternFill("solid", fgColor="B6D7A8")    # verde más fuerte

    df_temp = df.copy()
    df_temp["MES_AÑO"] = df_temp["FECHAVTO"].dt.strftime("%b-%y")

    pivot = df_temp.pivot_table(
        index="PROVEEDOR",
        columns="MES_AÑO",
        values="SALDO",
        aggfunc="sum",
        fill_value=0
    )
    columnas_ordenadas = sorted(
    [c for c in pivot.columns],
    key=lambda x: pd.to_datetime(x, format="%b-%y")
    )
    pivot = pivot[columnas_ordenadas]

    pivot["SALDO TOTAL"] = pivot.sum(axis=1)

    row = 3

    # Encabezados
    ws3.cell(row=row, column=1, value="PROVEEDOR").font = bold
    ws3.cell(row=row, column=1).fill = fill_header
    ws3.cell(row=row, column=1).border = border

    col = 2
    for mes in pivot.columns:
        c = ws3.cell(row=row, column=col, value=mes)
        c.font = bold
        c.alignment = right
        c.fill = fill_header
        c.border = border
        col += 1

    row += 1
    data_start = row

    # Datos
    for proveedor, data in pivot.iterrows():
        ws3.cell(row=row, column=1, value=proveedor).border = border
        ws3.cell(row=row, column=1).fill = fill_body

        col = 2
        for val in data:
            c = ws3.cell(row=row, column=col, value=val)
            c.alignment = right
            c.border = border

            # Colorear total
            if col == len(data) + 1:
                c.fill = fill_total
            else:
                c.fill = fill_body

            col += 1
        row += 1

    # Ajustar anchos
    for c in range(1, col):
        ws3.column_dimensions[get_column_letter(c)].width = 22

# ======================================================
# HOJA 4 – AGING REPORT
# ======================================================
    ws4 = wb.create_sheet(title="Aging Report")

    ws4.cell(row=1, column=1, value="Aging Report – Antigüedad de Saldos").font = bold_big

    headers = [
        "PROVEEDOR",
        # Rangos por días
        "0-30 días",
        "31-60 días",
        "61-90 días",
        "+90 días",
        "SALDO TOTAL (días)",
        # Rangos por semanas
        "Sem 1 (0-7d)",
        "Sem 2 (8-14d)",
        "Sem 3 (15-21d)",
        "Sem 4 (22-28d)",
        "+4 Sem (+28d)",
        "SALDO TOTAL (sem)",
    ]

    # Colores
    fill_header     = PatternFill("solid", fgColor="D9EAD3")  # verde claro
    fill_30         = PatternFill("solid", fgColor="E2F0D9")  # verde suave
    fill_60         = PatternFill("solid", fgColor="FFF2CC")  # amarillo
    fill_90         = PatternFill("solid", fgColor="FCE5CD")  # naranja
    fill_90p        = PatternFill("solid", fgColor="F4CCCC")  # rojo
    fill_total      = PatternFill("solid", fgColor="B6D7A8")  # verde fuerte
    # Colores semanas (azules)
    fill_sem1       = PatternFill("solid", fgColor="DAEEF3")  # azul muy suave
    fill_sem2       = PatternFill("solid", fgColor="BDD7EE")  # azul suave
    fill_sem3       = PatternFill("solid", fgColor="9DC3E6")  # azul medio
    fill_sem4       = PatternFill("solid", fgColor="6FA8DC")  # azul
    fill_sem_plus   = PatternFill("solid", fgColor="4472C4")  # azul fuerte
    fill_total_sem  = PatternFill("solid", fgColor="1F4E79")  # azul oscuro

    row = 3

    for col, h in enumerate(headers, start=1):
        c = ws4.cell(row=row, column=col, value=h)
        c.font = bold
        c.fill = fill_header
        c.border = border
        c.alignment = right if col > 1 else None

    row += 1

    hoy = pd.Timestamp.today().normalize()

    df_aging = df.copy()
    df_aging["DIAS_VENCIDO"] = (hoy - df_aging["FECHAVTO"]).dt.days

    def rango_dias(dias):
        if dias <= 30:  return "0-30"
        elif dias <= 60: return "31-60"
        elif dias <= 90: return "61-90"
        else:            return "+90"

    def rango_semanas(dias):
        if dias <= 7:   return "0-7"
        elif dias <= 14: return "8-14"
        elif dias <= 21: return "15-21"
        elif dias <= 28: return "22-28"
        else:            return "+28"

    df_aging["RANGO_DIAS"] = df_aging["DIAS_VENCIDO"].apply(rango_dias)
    df_aging["RANGO_SEM"]  = df_aging["DIAS_VENCIDO"].apply(rango_semanas)

    pivot_dias = df_aging.pivot_table(
        index="PROVEEDOR", columns="RANGO_DIAS",
        values="SALDO", aggfunc="sum", fill_value=0
    )
    pivot_sem = df_aging.pivot_table(
        index="PROVEEDOR", columns="RANGO_SEM",
        values="SALDO", aggfunc="sum", fill_value=0
    )

    for col in ["0-30", "31-60", "61-90", "+90"]:
        if col not in pivot_dias.columns: pivot_dias[col] = 0

    for col in ["0-7", "8-14", "15-21", "22-28", "+28"]:
        if col not in pivot_sem.columns:  pivot_sem[col]  = 0

    pivot_dias = pivot_dias[["0-30", "31-60", "61-90", "+90"]]
    pivot_sem  = pivot_sem[["0-7", "8-14", "15-21", "22-28", "+28"]]

    pivot_dias["TOTAL_DIAS"] = pivot_dias.sum(axis=1)
    pivot_sem["TOTAL_SEM"]   = pivot_sem.sum(axis=1)

    pivot_full = pivot_dias.join(pivot_sem, how="outer").fillna(0)

    col_fills = {
        2: fill_30,
        3: fill_60,
        4: fill_90,
        5: fill_90p,
        6: fill_total,
        7: fill_sem1,
        8: fill_sem2,
        9: fill_sem3,
        10: fill_sem4,
        11: fill_sem_plus,
        12: fill_total_sem,
    }

    dark_cols = {12}

    for proveedor, data in pivot_full.iterrows():
        ws4.cell(row=row, column=1, value=proveedor).border = border

        for col_idx, val in enumerate(data, start=2):
            c = ws4.cell(row=row, column=col_idx, value=val)
            c.alignment = right
            c.border = border
            c.fill = col_fills.get(col_idx, fill_30)
            if col_idx in dark_cols:
                c.font = Font(color="FFFFFF")

        row += 1

    ws4.column_dimensions["A"].width = 35
    for i, col_letter in enumerate(["B","C","D","E","F","G","H","I","J","K","L"], start=1):
        ws4.column_dimensions[col_letter].width = 18
    
# ======================================================
# HOJA 5 – TOP 10 PROVEEDORES
# ======================================================
    ws5 = wb.create_sheet(title="Top 10 Proveedores")

    ws5.cell(row=1, column=1, value="Top 10 Proveedores por Saldo").font = bold_big

    headers = ["PROVEEDOR", "SALDO TOTAL"]

    # Colores
    fill_header = PatternFill("solid", fgColor="D9D2E9")   # morado claro
    fill_top1 = PatternFill("solid", fgColor="FFD966")    # oro
    fill_top2 = PatternFill("solid", fgColor="D9D9D9")    # plata
    fill_top3 = PatternFill("solid", fgColor="F4B183")    # bronce
    fill_normal = PatternFill("solid", fgColor="EFEFEF") # gris claro
    
    row = 3

    for col, h in enumerate(headers, start=1):
        c = ws5.cell(row=row, column=col, value=h)
        c.font = bold
        c.fill = fill_header
        c.border = border
        c.alignment = right if col == 2 else None

    row += 1

    top10 = (
        df.groupby("PROVEEDOR")["SALDO"]
        .sum()
        .sort_values(ascending=False)
        .head(10)
        .reset_index()
    )

    rank = 1
    for _, r in top10.iterrows():
        ws5.cell(row=row, column=1, value=r["PROVEEDOR"])
        ws5.cell(row=row, column=1).border = border

        c = ws5.cell(row=row, column=2, value=r["SALDO"])
        c.alignment = right
        c.border = border

        # Colores por ranking
        if rank == 1:
            fill = fill_top1
        elif rank == 2:
            fill = fill_top2
        elif rank == 3:
            fill = fill_top3
        else:
            fill = fill_normal

        ws5.cell(row=row, column=1).fill = fill
        ws5.cell(row=row, column=2).fill = fill

        rank += 1
        row += 1


    ws5.column_dimensions["A"].width = 40
    ws5.column_dimensions["B"].width = 20

# ======================================================
# HOJA 6 – TABLAS POR PROVEEDOR
# ======================================================
    ws6 = wb.create_sheet(title="Detalle por proveedor")

    ws6.cell(row=1, column=1, value="Detalle Mensual por Proveedor").font = bold_big

    row = 3
    tabla_counter = 1

    # Colores
    fill_header = PatternFill("solid", fgColor="BDD7EE")   # azul claro
    fill_total = PatternFill("solid", fgColor="D9EAD3")    # verde claro

    for proveedor in sorted(df["PROVEEDOR"].unique()):
        sub = df[df["PROVEEDOR"] == proveedor]

        # Título proveedor
        ws6.cell(row=row, column=1, value=proveedor).font = bold_big
        row += 2

        # Encabezados
        headers = ["MES", "AÑO", "SALDO DEL MES"]
        header_row = row

        for col, h in enumerate(headers, start=1):
            c = ws6.cell(row=row, column=col, value=h)
            c.font = bold
            c.fill = fill_header
            c.border = border
            if col == 3:
                c.alignment = right

        row += 1
        start_data_row = row

        resumen = (
            sub.groupby(["MES", "AÑO"])["SALDO"]
            .sum()
            .reset_index()
            .sort_values(by=["AÑO"])
        )

        total_proveedor = 0

        for _, r in resumen.iterrows():
            ws6.cell(row=row, column=1, value=r["MES"]).border = border
            ws6.cell(row=row, column=2, value=r["AÑO"]).border = border

            c = ws6.cell(row=row, column=3, value=r["SALDO"])
            c.alignment = right
            c.border = border

            total_proveedor += r["SALDO"]
            row += 1

        # Total proveedor (fuera de la tabla)
        ws6.cell(row=row, column=1, value="TOTAL PROVEEDOR").font = bold
        ws6.cell(row=row, column=1).fill = fill_total
        ws6.cell(row=row, column=1).border = border

        c = ws6.cell(row=row, column=3, value=total_proveedor)
        c.font = bold
        c.alignment = right
        c.fill = fill_total
        c.border = border

        row += 3  # espacio entre proveedores

    ws6.column_dimensions["A"].width = 25
    ws6.column_dimensions["B"].width = 15
    ws6.column_dimensions["C"].width = 22

    wb.save(output_path)

# -----------------------------
# 3. Ejecutar reporte
# -----------------------------
def run_report(
    input_path: str,
    output_dir: str = "output",
    base_name: str = "ReporteAnaliticoCXP"
):
    os.makedirs(output_dir, exist_ok=True)
    today = datetime.today().strftime("%Y%m%d")

    output_path = os.path.join(
        output_dir,
        f"{base_name}{today}.xlsx"
    )
  
    df = load_and_clean(input_path)
    generate_xlsx(df, output_path)

    print(f"✅ Archivo generado: {output_path}")